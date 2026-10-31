"""Seed the project-kept reference data from last season's sheet (T029, FR-005, FR-040).

Run from the repository root:

    uv run --with openpyxl python -I tools/reference-seed/seed.py

Writes:
  app/src/reference/shared-lists.seed.json  the 11 Sabhas and the Artists in the sheet
  app/src/reference/composer-dates.json     composer key -> born, died, approx
  app/src/reference/topics.json             subject key -> Topic
  app/src/reference/forms.json              the project-kept forms with no Composition
  tools/reference-seed/unmatched.md         what still needs the Maintainer

Composer dates and topics are keyed by catalog natural keys (the composer `name_key`, the
subject name). Those keys come from the catalog snapshot, which does not exist yet (T021), so
both files stay empty and every name is listed in unmatched.md instead of being guessed.
"""

import difflib
import json
import re
import uuid
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SHEET = ROOT / "docs/reference/wrapped-2024-25.xlsx"
OUT = ROOT / "app/src/reference"
REPORT = ROOT / "tools/reference-seed/unmatched.md"

# Shared ids stay the same on every run, so re-seeding never makes a second copy of an entry.
NAMESPACE = uuid.UUID("6f1c2a52-3b0e-4f43-9d1e-7a4c0b8e5d21")
LIST_VERSION = 1

# Project-kept forms an Item uses when it has no Composition (data-model.md, Item `form`).
# Where the catalog has the same form, the key is the catalog's ("viruttam", "shlokam"), so
# the two never split one form's count.
FORMS = [
    {"key": "rtp", "display": "RTP", "aliases": ["Ragam Tanam Pallavi", "Ragam Thanam Pallavi"]},
    {"key": "viruttam", "display": "Viruttam", "aliases": ["Virutham", "Viruttham"]},
    {"key": "shlokam", "display": "Slokam", "aliases": ["Sloka", "Shlokam", "Shloka"]},
    {"key": "medley", "display": "Medley", "aliases": ["Mix"]},
]


def clean(value):
    """A cell as trimmed text, or None when it is empty."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def rows(sheet):
    """The data rows of a sheet as dicts keyed by header, skipping blank rows.

    A repeated header gets a number ("Presenter/Co-presenter 2"), so no column is lost: the
    Lecdems sheet has three columns with the same name.
    """
    values = list(sheet.iter_rows(values_only=True))
    seen = Counter()
    header = []
    for cell in values[0]:
        name = clean(cell) or ""
        seen[name] += 1
        header.append(name if seen[name] == 1 else f"{name} {seen[name]}")
    for row in values[1:]:
        if any(clean(cell) for cell in row):
            yield {header[i]: row[i] for i in range(min(len(header), len(row)))}


def shared(kind, names):
    return [
        {
            "id": str(uuid.uuid5(NAMESPACE, f"{kind}:{name}")),
            "name": name,
            "aliases": [],
            "listVersion": LIST_VERSION,
        }
        for name in names
    ]


def near_duplicates(names):
    """Pairs of names that look like one person spelled twice, for the Maintainer to judge."""
    key = lambda name: re.sub(r"[^a-z]", "", name.lower())
    pairs = []
    for i, a in enumerate(names):
        for b in names[i + 1 :]:
            if difflib.SequenceMatcher(None, key(a), key(b)).ratio() >= 0.85:
                pairs.append((a, b))
    return pairs


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def main():
    book = openpyxl.load_workbook(SHEET, read_only=True, data_only=True)

    concerts = list(rows(book["Main"]))
    lecdems = list(rows(book["Lecdems"]))
    artist_columns = [c for c in concerts[0] if c.startswith(("Artist", "Accompaniment"))]

    sabhas = list(dict.fromkeys(clean(row["Sabha"]) for row in concerts if clean(row["Sabha"])))
    artists = list(
        dict.fromkeys(
            name
            for row in concerts
            for column in artist_columns
            if (name := clean(row.get(column)))
        )
    )
    for row in lecdems:
        for column, value in row.items():
            if column.startswith("Presenter") and (name := clean(value)) and name not in artists:
                artists.append(name)

    composers = []
    for row in book["Timelines"].iter_rows(values_only=True):
        name = clean(row[0]) if row else None
        if not name:
            continue
        born = row[1] if len(row) > 1 else None
        died = row[2] if len(row) > 2 else None
        composers.append(
            {
                "name": name,
                "born": int(born) if born is not None else None,
                "died": int(died) if died is not None else None,
            }
        )

    topics = Counter(
        topic
        for row in rows(book["Compositions"])
        if (topic := clean(row.get("Topic"))) and topic not in {"?", "-", "None"}
    )

    write_json(OUT / "shared-lists.seed.json", {"sabhas": shared("sabha", sabhas), "artists": shared("artist", artists)})
    write_json(OUT / "composer-dates.json", {})
    write_json(OUT / "topics.json", {})
    write_json(OUT / "forms.json", FORMS)

    lines = [
        "# Reference seed: names for the Maintainer",
        "",
        "Written by `tools/reference-seed/seed.py` from `docs/reference/wrapped-2024-25.xlsx`.",
        "Nothing here is guessed: each list waits on a decision or on the catalog snapshot.",
        "",
        f"## Composers not yet keyed ({len(composers)})",
        "",
        "`composer-dates.json` is keyed by the composer `name_key` from the catalog snapshot,",
        "which does not exist yet (T021). Once it does, each name below is matched to a key,",
        "and the ones that don't match stay listed here. `approx` is not in the sheet; it is",
        "decided per composer then.",
        "",
        "| Composer (as in the sheet) | Born | Died |",
        "|---|---|---|",
        *(
            f"| {c['name']} | {c['born'] if c['born'] is not None else ''} | {c['died'] if c['died'] is not None else ''} |"
            for c in composers
        ),
        "",
        f"## Topics not yet keyed ({len(topics)})",
        "",
        "`topics.json` maps a catalog Subject to a Topic. The sheet records only the Topic, so",
        "each one is paired with the catalog Subjects of the Compositions it was given to once",
        "the snapshot exists. Counts are Items in the sheet.",
        "",
        "| Topic | Items |",
        "|---|---|",
        *(f"| {topic} | {count} |" for topic, count in sorted(topics.items(), key=lambda t: (-t[1], t[0]))),
        "",
        "## Artist names that may be one person",
        "",
        "Seeded as separate Shared Artists. Merge them, or keep one as an alias, if they are the",
        "same person.",
        "",
        *(f"- {a} / {b}" for a, b in near_duplicates(artists)),
        "",
        f"Seeded: {len(sabhas)} Sabhas, {len(artists)} Artists, {len(FORMS)} forms.",
        "",
    ]
    REPORT.write_text("\n".join(lines), encoding="utf-8")
    print(f"{len(sabhas)} Sabhas, {len(artists)} Artists, {len(composers)} composers listed, {len(topics)} topics listed")


if __name__ == "__main__":
    main()
